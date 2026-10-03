#!/usr/bin/env python3
"""Build the vmbdiet two-cohort validation deck (docs/vmbdiet-validation.pptx)."""
from __future__ import annotations
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.join(ROOT, "docs", "vmbdiet-validation.pptx")

INK   = RGBColor(0x1A, 0x1A, 0x1A)
MUTE  = RGBColor(0x5A, 0x5A, 0x5A)
BAR   = RGBColor(0x0F, 0x3A, 0x33)   # deep teal-green header
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1B, 0x78, 0x37)   # replicate / crispatus / good
TEAL  = RGBColor(0x0F, 0x6B, 0x5F)   # accent
AMBER = RGBColor(0xC0, 0x8A, 0x2D)   # iners
RED   = RGBColor(0xB2, 0x3A, 0x5B)   # CST-IV / BV
BLUE  = RGBColor(0x5A, 0x7D, 0x8C)   # jensenii
PURP  = RGBColor(0x7A, 0x6A, 0xA8)   # gasseri
CREAM = RGBColor(0xF6, 0xF4, 0xEF)
PANEL = RGBColor(0xFB, 0xFA, 0xF5)
LINE  = RGBColor(0xDD, 0xD6, 0xC8)
MINT  = RGBColor(0x9F, 0xD3, 0xB8)

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s, x, y, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp
def rrect(s, x, y, w, h, color, line=LINE):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.color.rgb = line; sp.line.width = Pt(0.75); sp.shadow.inherit = False
    return sp
def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return tf
def para(tf, text, size, color=INK, bold=False, first=False, align=PP_ALIGN.LEFT,
         space_after=6, bullet=False, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
    runs = text if isinstance(text, list) else [(text, color, bold, italic)]
    for i, item in enumerate(runs):
        t, c, b, it = item if isinstance(item, tuple) else (item, color, bold, italic)
        r = p.add_run(); r.text = (("•  " if bullet else "") + t) if i == 0 else t
        r.font.size = Pt(size); r.font.color.rgb = c; r.font.bold = b; r.font.italic = it
        r.font.name = "Calibri"
    return p
def header(s, title, kicker=None):
    rect(s, 0, 0, SW, Inches(1.1), BAR)
    tf = box(s, Inches(0.6), Inches(0.14), Inches(12.1), Inches(0.85), MSO_ANCHOR.MIDDLE)
    if kicker:
        para(tf, kicker, 12, MINT, bold=True, first=True, space_after=1)
        para(tf, title, 25, WHITE, bold=True)
    else:
        para(tf, title, 26, WHITE, bold=True, first=True)
def statcard(s, x, y, w, lab, val, accent=GREEN, h=Inches(1.25)):
    rrect(s, x, y, w, h, PANEL)
    rect(s, x, y, Inches(0.08), h, accent)
    tf = box(s, x + Inches(0.28), y + Inches(0.14), w - Inches(0.45), h - Inches(0.28), MSO_ANCHOR.MIDDLE)
    para(tf, lab, 13, MUTE, bold=True, first=True, space_after=5)
    para(tf, val, 17, INK, bold=True)

# 1 ---- title
s = slide()
rect(s, 0, 0, SW, SH, BAR)
rect(s, 0, Inches(4.7), SW, Inches(0.06), GREEN)
tf = box(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(2.8))
para(tf, "Diet & the vaginal microbiome", 44, WHITE, bold=True, first=True, space_after=6)
para(tf, "Validating a reusable diet → community-state-type analysis toolkit", 22, RGBColor(0xD8,0xE8,0xDF), space_after=2)
para(tf, "on two independent public cohorts, reprocessed from raw reads", 22, RGBColor(0xD8,0xE8,0xDF))
tf = box(s, Inches(0.9), Inches(5.05), Inches(11.5), Inches(1.6))
para(tf, "Jean Yu   ·   October 2026", 16, RGBColor(0xBF,0xD6,0xC8), bold=True, first=True, space_after=3)
para(tf, "github.com/jyu9675/vmbdiet", 14, MINT)

# 2 ---- the question
s = slide(); header(s, "The question", "motivation")
tf = box(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(1.5))
para(tf, "Diet is reported to shape the vaginal microbiome — but every study uses its own ad-hoc "
         "pipeline, so the signal is hard to compare or build on.", 18, INK, first=True, space_after=10)
para(tf, [("I packaged the shared analysis (CST classification + covariate-adjusted diet association) into one "
           "tested tool, ", INK, False, False), ("vmbdiet", TEAL, True, False),
          (", then asked: does it recover a published result when run, from raw data, on a cohort it "
           "was not tuned to?", INK, False, False)], 18)
for i,(txt) in enumerate([
    "Neggers 2007 — dietary fat ↑ BV; folate / vit-A / calcium ↓ BV",
    "Tuddenham 2019 — dairy / fiber / vit-D → L. crispatus over L. iners",
    "Hawaii 2024 — carbohydrate / HEI-2015 → L. crispatus",
    "Italy 2025 — animal protein → CST-IV; alcohol → Gardnerella"]):
    y = Inches(3.5 + i*0.62)
    rrect(s, Inches(0.6), y, Inches(12.1), Inches(0.5), PANEL)
    tf = box(s, Inches(0.85), y, Inches(11.7), Inches(0.5), MSO_ANCHOR.MIDDLE)
    para(tf, txt, 15, INK, first=True)
tf = box(s, Inches(0.6), Inches(6.08), Inches(12.1), Inches(0.5))
para(tf, "Four studies, one shared analysis — standardized and validated here.", 14, MUTE, italic=True, first=True)

# 3 ---- approach
s = slide(); header(s, "Approach", "two public cohorts · from raw reads · one pipeline")
steps = ["raw FASTQ (SRA)","cutadapt\nprimers","vsearch\nmerge","EE<1.0\nfilter","UNOISE3\nASVs",
         "RDP + RefSeq\ntaxonomy","vmbdiet\nCST + assoc"]
x = Inches(0.55); w = Inches(1.62); y = Inches(1.6); gap = Inches(0.17)
for i,st in enumerate(steps):
    rrect(s, x, y, w, Inches(1.0), PANEL if i else CREAM, line=LINE)
    tf = box(s, x+Inches(0.08), y, w-Inches(0.16), Inches(1.0), MSO_ANCHOR.MIDDLE)
    para(tf, st, 12, INK, bold=(i in(0,6)), first=True, align=PP_ALIGN.CENTER)
    if i < len(steps)-1:
        ar = box(s, x+w, y, gap, Inches(1.0), MSO_ANCHOR.MIDDLE)
        para(ar, "→", 16, TEAL, bold=True, first=True, align=PP_ALIGN.CENTER)
    x = x + w + gap
tf = box(s, Inches(0.6), Inches(3.1), Inches(12.1), Inches(1.0))
para(tf, [("Nothing reuses the authors' processed tables ", INK, True, False),
          ("— the ASV table is rebuilt from raw reads; CSTs by dominant-taxon rule (VALENCIA centroids "
           "shipped); associations by per-SD adjusted logistic + BH-FDR, Spearman, and PERMANOVA. Runs on a laptop.",
           INK, False, False)], 16, first=True)
statcard(s, Inches(0.6), Inches(4.5), Inches(3.9), "Cohort 1 · Italian “StressDiet”",
         "N=113 · non-pregnant · FFQ", TEAL, h=Inches(1.5))
statcard(s, Inches(4.7), Inches(4.5), Inches(3.9), "Cohort 2 · MELODY",
         "N=47 · pregnancy±IBD · HEI-2015", TEAL, h=Inches(1.5))
statcard(s, Inches(8.8), Inches(4.5), Inches(3.9), "Dropped · MicrobeMom",
         "gut/stool, ~1.5 TB — no vaginal", RED, h=Inches(1.5))
tf = box(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.5))
para(tf, "MicrobeMom (PRJEB48251) was verified at the ENA record to be a maternal-gut cohort with no vaginal "
         "samples — correctly excluded, not run.", 13, MUTE, italic=True, first=True)

# 4 ---- cohort 1 communities
s = slide(); header(s, "Cohort 1 — the communities look right", "Italian StressDiet · N=113")
tf = box(s, Inches(0.6), Inches(1.3), Inches(12.1), Inches(0.6))
para(tf, "Before any diet test means anything, reconstructed communities must resemble a real cohort. "
         "They do (species-level resolution worked).", 16, INK, first=True)
# CST stacked bar
cst = [("CST I  (L. crispatus)",49,GREEN),("CST III  (L. iners)",22,AMBER),
       ("CST IV  (diverse/BV)",35,RED),("CST V",4,BLUE),("CST II",3,PURP)]
total = sum(c[1] for c in cst)
x = Inches(0.6); y = Inches(2.3); full = Inches(12.1); barh = Inches(0.8)
for lab,n,col in cst:
    w = int(full * n/total)
    seg = rect(s, x, y, w, barh, col)
    if n/total > 0.07:
        tf = box(s, x, y, w, barh, MSO_ANCHOR.MIDDLE)
        para(tf, f"{n}", 15, WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
    x = x + w
# legend
lx = Inches(0.6); ly = Inches(3.35)
for lab,n,col in cst:
    rect(s, lx, ly+Inches(0.04), Inches(0.22), Inches(0.22), col)
    tf = box(s, lx+Inches(0.3), ly, Inches(2.7), Inches(0.35), MSO_ANCHOR.MIDDLE)
    para(tf, lab, 12, INK, first=True)
    lx = lx + Inches(2.55)
tf = box(s, Inches(0.6), Inches(4.2), Inches(12.1), Inches(0.5))
para(tf, "Top taxa by mean relative abundance:", 15, INK, bold=True, first=True)
taxa = [("L. crispatus","41%"),("L. iners","22%"),("G. vaginalis","6.3% (81% prev.)"),
        ("Prevotella","4.8%"),("L. gasseri","3.9%"),("L. jensenii","3.7%")]
x = Inches(0.6)
for nm,v in taxa:
    rrect(s, x, Inches(4.75), Inches(1.93), Inches(1.0), PANEL)
    tf = box(s, x+Inches(0.12), Inches(4.85), Inches(1.7), Inches(0.8), MSO_ANCHOR.MIDDLE)
    para(tf, nm, 12.5, INK, bold=True, first=True, space_after=3, align=PP_ALIGN.CENTER)
    para(tf, v, 13, TEAL, bold=True, align=PP_ALIGN.CENTER)
    x = x + Inches(2.02)
tf = box(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.5))
para(tf, "5.1M merged reads → 1,312 ASVs → 163 taxa; all 113 samples retained.", 13, MUTE, italic=True, first=True)

# 5 ---- cohort 1 results
s = slide(); header(s, "Cohort 1 — both published signals replicate", "Italian StressDiet · independent reprocessing")
statcard(s, Inches(0.6), Inches(1.4), Inches(5.95), "Animal-protein fraction → CST-IV",
         "OR 1.66 / SD  (1.09–2.53)  p = 0.019", GREEN, h=Inches(1.4))
statcard(s, Inches(6.75), Inches(1.4), Inches(5.95), "Alcohol → Gardnerella / anaerobes",
         "ρ +0.30, p=0.001  ·  Leptotrichia q=0.012", GREEN, h=Inches(1.4))
rrect(s, Inches(0.6), Inches(3.15), Inches(12.1), Inches(2.0), CREAM, line=AMBER)
tf = box(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(1.7))
para(tf, "The decisive methodological step: energy adjustment", 16, AMBER, bold=True, first=True, space_after=8)
para(tf, "On raw intake, fiber and vegetable protein looked protective — but that was a confound of lower "
         "total energy in CST-IV women (1,754 vs 1,995 kcal, p=0.016). Expressing nutrients as density per 1,000 "
         "kcal removed the artifact and left the animal-protein effect standing.", 15, INK, space_after=6)
para(tf, "The tool handled a standard nutritional-epidemiology pitfall correctly — and the association that "
         "survives is the one the original authors reported.", 15, INK, italic=True)
tf = box(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(0.9))
para(tf, "PERMANOVA (community ~ CST-IV) pseudo-F ≈ 22, p = 0.001 — a structural sanity check "
         "(partly definitional), not a diet result.", 14, MUTE, first=True)

# 6 ---- cohort 2 MELODY
s = slide(); header(s, "Cohort 2 — an honest null (and why that's reassuring)", "MELODY · pregnancy±IBD · N=47")
# comparison table
cols = ["Comparable test","Italian (N=113)","MELODY (N=47)"]
rows = [("Animal protein → CST-IV","OR 1.66, p=0.019","OR 1.30, p=0.44"),
        ("HEI-2015 → CST-IV","—","OR 1.40, p=0.33"),
        ("Alcohol → Gardnerella","ρ +0.30, p=0.001","≈0 intake (untestable)")]
x0=Inches(0.6); y0=Inches(1.4); cw=[Inches(5.0),Inches(3.55),Inches(3.55)]; rh=Inches(0.62)
cx=x0
for j,c in enumerate(cols):
    rect(s, cx, y0, cw[j], rh, BAR)
    tf=box(s, cx+Inches(0.12), y0, cw[j]-Inches(0.2), rh, MSO_ANCHOR.MIDDLE)
    para(tf, c, 13, WHITE, bold=True, first=True)
    cx=cx+cw[j]
for i,r in enumerate(rows):
    cx=x0; yy=y0+rh*(i+1)
    for j,val in enumerate(r):
        rect(s, cx, yy, cw[j], rh, PANEL if i%2==0 else CREAM, line=LINE)
        tf=box(s, cx+Inches(0.12), yy, cw[j]-Inches(0.2), rh, MSO_ANCHOR.MIDDLE)
        colr = GREEN if (j==1 and i in(0,2)) else INK
        para(tf, val, 13.5, colr, bold=(j==1 and i in(0,2)), first=True)
        cx=cx+cw[j]
rrect(s, Inches(0.6), Inches(3.9), Inches(12.1), Inches(2.4), CREAM, line=TEAL)
tf = box(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(2.1))
para(tf, "A null is the right answer here — and a good sign", 16, TEAL, bold=True, first=True, space_after=8)
para(tf, "Underpowered: 47 women, 12 CST-IV events — a per-SD OR ≈ 1.3 is undetectable.", 15, INK, bullet=True, space_after=5)
para(tf, "Pregnancy flattens the exposure: alcohol ≈ 0, and the community skews to Lactobacillus regardless of diet.", 15, INK, bullet=True, space_after=5)
para(tf, "The source study itself found no diet–microbiota difference — we reproduce that.", 15, INK, bullet=True, space_after=5)
para(tf, "The one comparable signal (animal protein) still points the same way. A tool manufacturing significance "
         "here would be the worry — not this.", 15, INK, italic=True)

# 7 ---- interpretation
s = slide(); header(s, "Interpretation & limits", "what this is — and isn't")
rrect(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(1.5), PANEL, line=GREEN)
tf = box(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.2), MSO_ANCHOR.MIDDLE)
para(tf, [("A positive replication in the powered cohort and an honest null in the underpowered one is exactly "
           "the behavior a trustworthy tool should show.", INK, True, False)], 17, first=True, space_after=5)
para(tf, "The animal-protein → CST-IV direction is consistent across both cohorts.", 15, MUTE)
tf = box(s, Inches(0.6), Inches(3.2), Inches(12.1), Inches(2.6))
para(tf, "Limits — stated plainly", 15, RED, bold=True, first=True, space_after=8)
for t in ["Both cohorts are modest and cross-sectional — no causal direction.",
          "The animal-protein association is nominally significant but does not clear FDR across all nutrients (q=0.14).",
          "Species resolution is limited to the dominant vaginal taxa; rare taxa are genus-level.",
          "PERMANOVA on CST-IV is a structural check, not a diet finding (CST-IV is itself a community label)."]:
    para(tf, t, 15, INK, bullet=True, space_after=7)

# 8 ---- what's next
s = slide(); header(s, "What this enables", "next steps")
items = [("Validated, reusable pipeline","The same code profiles any cohort pairing a 16S table with diet — a prospective or larger diet→CST study, or a multi-cohort meta-analysis, is now low-risk with the analysis in hand."),
         ("A third HEI-2015 cohort","The Hawaii multi-ethnic cohort (PMC11479099) has not yet deposited an accession — the ideal third replication once public. (Flagged to revisit.)"),
         ("Complements the strain program","Diet is a plausible modifier of which Lactobacillus strains establish and persist — links to the lab's strain-transmission / engraftment work.")]
y = Inches(1.45)
for t,d in items:
    rrect(s, Inches(0.6), y, Inches(12.1), Inches(1.45), PANEL)
    rect(s, Inches(0.6), y, Inches(0.08), Inches(1.45), TEAL)
    tf = box(s, Inches(0.95), y+Inches(0.16), Inches(11.5), Inches(1.15), MSO_ANCHOR.MIDDLE)
    para(tf, t, 17, TEAL, bold=True, first=True, space_after=5)
    para(tf, d, 14.5, INK)
    y = y + Inches(1.62)
rect(s, 0, Inches(6.95), SW, Inches(0.55), BAR)
tf = box(s, Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.55), MSO_ANCHOR.MIDDLE)
para(tf, [("Reproducible scripts, derived tables & write-up:  ", MINT, False, False),
          ("github.com/jyu9675/vmbdiet", WHITE, True, False)], 14, first=True)

prs.save(OUT)
print("wrote", OUT, "(%d slides)" % len(prs.slides._sldIdLst))
